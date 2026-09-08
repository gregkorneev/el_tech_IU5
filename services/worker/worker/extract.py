"""Small, page-aware extractors; binary source files never become persistent assets."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO


@dataclass(frozen=True)
class ExtractedPart:
    text: str
    page: int | None = None
    slide: int | None = None


def extract(name: str, content: bytes) -> list[ExtractedPart]:
    extension = name.rsplit(".", 1)[-1].lower()
    if extension in {"txt", "md"}:
        return [ExtractedPart(content.decode("utf-8", errors="replace"))]
    if extension == "pdf":
        from pypdf import PdfReader
        return [ExtractedPart(page.extract_text() or "", page=index + 1) for index, page in enumerate(PdfReader(BytesIO(content)).pages)]
    if extension == "docx":
        from docx import Document
        return [ExtractedPart("\n".join(p.text for p in Document(BytesIO(content)).paragraphs))]
    if extension == "pptx":
        from pptx import Presentation
        return [ExtractedPart("\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")), slide=index + 1) for index, slide in enumerate(Presentation(BytesIO(content)).slides)]
    if extension == "xlsx":
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        return [ExtractedPart("\n".join(" | ".join(str(cell) for cell in row if cell is not None) for row in sheet.iter_rows(values_only=True)), slide=index + 1) for index, sheet in enumerate(book.worksheets)]
    raise ValueError(f"No text extractor for .{extension}")


def chunks(parts: list[ExtractedPart], max_chars: int = 2_000) -> list[ExtractedPart]:
    result: list[ExtractedPart] = []
    for part in parts:
        words, current = part.text.split(), []
        for word in words:
            if current and len(" ".join(current)) + len(word) + 1 > max_chars:
                result.append(ExtractedPart(" ".join(current), part.page, part.slide)); current = []
            current.append(word)
        if current: result.append(ExtractedPart(" ".join(current), part.page, part.slide))
    return result
